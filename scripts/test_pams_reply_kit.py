# -*- coding: utf-8 -*-
"""pams_reply_kit 시험 — python scripts/test_pams_reply_kit.py (게이트는 npx tsx 로 실제 실행)."""
import json
import os
import sys
import tempfile
import unittest

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
import pams_reply_kit as rk  # noqa: E402

LINK = "개인마다 계약이 달라 여기서 말씀드리기 어려워요. 프로필 링크로 상담 남겨 주시면 직접 보고 말씀드릴게요."
GOOD = [LINK, "읽어 주셔서 고맙습니다.", "해지는 되돌리기 어려워서, 담보별로 펼쳐 보고 판단하시는 게 안전해요. 바로 해지하지 마시고요.",
        "마음이 많이 쓰이시겠어요. 어머님 빨리 회복하시길 바랄게요."]
BAD = ["○○사 간병보험 추천해요. 제일 좋아요.", "보통 월 5만원대면 돼요.", "감사합니다! 무료 상담도 해드려요~",
       "네 전화드릴게요! (번호 되받기)"]
POST = "https://www.threads.com/@goodfinance_sj/post/Dd3Le5YASPf"


class T(unittest.TestCase):
    def test_good_to_sheet_bad_to_blocked(self):
        from openpyxl import load_workbook
        items = [{"post_url": POST, "comment": "부모님 거 봐주세요 010-1234-5678", "reply": r} for r in GOOD + BAD]
        with tempfile.TemporaryDirectory() as d:
            s = rk.build(items, out_dir=d, date="0930")
            self.assertEqual((s["접수"], s["막힘"]), (len(GOOD), len(BAD)))
            self.assertTrue(s["xlsx"].endswith(f"0930_댓글_{len(GOOD)}건.xlsx"))
            wb = load_workbook(s["xlsx"])
            rows = list(wb["접수"].iter_rows(values_only=True))
            self.assertEqual(rows[0][0], "번호")
            self.assertEqual([r[4] for r in rows[1:]], GOOD, "통과분만, 원고 그대로")
            self.assertTrue(all("****" in r[3] and "1234-5678" not in r[3] for r in rows[1:]), "댓글 전화번호는 가린다")
            self.assertEqual(len(list(wb["막힘"].iter_rows(values_only=True))) - 1, len(BAD))

    def test_from_replies_filters(self):
        d = {"댓글": {
            "1": {"계정": "goodfinance", "초안": "읽어 주셔서 고맙습니다.", "상태": "승인", "글": "잘 봤어요", "뿌리글_짧은": "AAA", "링크": "L1"},
            "2": {"계정": "goodfinance", "초안": None, "상태": "승인", "글": "x"},
            "3": {"계정": "baksatravel", "초안": "안녕하세요", "상태": "승인", "글": "y"},
            "4": {"계정": "goodfinance", "초안": "고맙습니다", "상태": "초안", "글": "z"},
        }}
        with tempfile.NamedTemporaryFile("w", suffix=".json", delete=False, encoding="utf-8") as fp:
            json.dump(d, fp, ensure_ascii=False)
        try:
            items = rk.from_replies(fp.name, "승인")
            default = rk.from_replies(fp.name)
        finally:
            os.unlink(fp.name)
        self.assertEqual(len(items), 1)
        self.assertEqual(items[0]["post_url"], "https://www.threads.com/@goodfinance_sj/post/AAA")
        self.assertEqual([x["id"] for x in default], ["4"], "기본은 reply_bot 정본 상태 「초안」")

    def test_empty_reply_stops(self):
        with self.assertRaises(rk.KitError):
            rk.build([{"reply": "  "}], dry_run=True)


class Gate(unittest.TestCase):
    """🔴 답글 게시 게이트 — 승인 · 심의필 · 유효기간 · 필수안내 · 해시 일치가 모두 맞아야 보낸다."""
    TEXT = "읽어 주셔서 고맙습니다."

    def row(self, **k):
        r = {"status": "승인", "review_no": "2026-10-1234", "review_from": "2026-10-01", "review_to": "2027-09-30",
             "notice_text": "상기 내용은 보험설계사의 의견이며…", "reply_text": self.TEXT, "reply_sha256": rk.reply_hash(self.TEXT)}
        r.update(k)
        return r

    def test_all_good_sends(self):
        self.assertEqual(rk.posting_verdict(self.row(), self.TEXT, "2026-10-02"), (True, ""))

    def test_each_missing_piece_blocks(self):
        cases = [(None, "기록이 없다"), (self.row(status="초안"), "승인 전"), (self.row(status="심의대기"), "승인 전"),
                 (self.row(status="반송"), "승인 전"), (self.row(review_no="제2026-10-1234호"), "형식"),
                 (self.row(review_no=None), "형식"), (self.row(review_to="2026-10-01"), "유효기간 밖"),
                 (self.row(notice_text=" "), "필수안내"), (self.row(status="답함"), "이미 보냈다"),
                 (self.row(posted_reply_id="R1"), "이미 보냈다")]
        for row, why in cases:
            ok, msg = rk.posting_verdict(row, self.TEXT, "2026-10-02")
            self.assertFalse(ok, row)
            self.assertIn(why, msg)

    def test_one_character_change_blocks(self):
        for t in (self.TEXT + " ", self.TEXT.replace(".", "!"), self.TEXT + "\n"):
            self.assertEqual(rk.posting_verdict(self.row(), t, "2026-10-02"), (False, "심의본과 다르다 — 한 글자라도 바뀌면 신규 심의"))

    def test_hash_is_utf8_exact(self):
        import hashlib
        self.assertEqual(rk.reply_hash("가"), hashlib.sha256("가".encode("utf-8")).hexdigest())
        self.assertNotEqual(rk.reply_hash("가\r\n"), rk.reply_hash("가\n"))

    def test_compose_post_appends_notice_and_guards_length(self):
        self.assertEqual(rk.compose_post(self.row()), self.TEXT + "\n\n상기 내용은 보험설계사의 의견이며…")
        with self.assertRaises(rk.KitError):
            rk.compose_post(self.row(reply_text="가" * 400, notice_text="나" * 120))

    def test_statuses_match_migration(self):
        s = open(os.path.join(os.path.dirname(os.path.dirname(os.path.abspath(__file__))), "sql", "005_reply_reviews.sql"), encoding="utf-8").read()
        self.assertIn("check (status in (" + ", ".join(f"'{x}'" for x in rk.STATUSES) + "))", s)


class Record(unittest.TestCase):
    def test_rows_and_missing_table(self):
        items = [{"id": "C1", "root_post_id": "P1", "post_url": "u", "reply": "고맙습니다.", "comment_url": "c"},
                 {"reply": "id 없는 --input 행"}]
        rows = rk.record_rows(items)
        self.assertEqual([(r["comment_id"], r["status"]) for r in rows], [("C1", "초안")])

        class Resp:
            def __init__(self, code, text=""):
                self.status_code, self.text = code, text
        self.assertEqual(rk.record(items, post=lambda rs: Resp(404, '{"code":"PGRST205"}'))[0], 0)
        self.assertEqual(rk.record(items, post=lambda rs: Resp(201))[0], 1)
        with self.assertRaises(rk.KitError):
            rk.record(items, post=lambda rs: Resp(400, "bad"))

    def test_sheet_has_hash_column(self):
        from openpyxl import load_workbook
        with tempfile.TemporaryDirectory() as d:
            s = rk.build([{"post_url": POST, "comment": "잘 봤어요", "reply": LINK}], out_dir=d, date="1001")
            rows = list(load_workbook(s["xlsx"])["접수"].iter_rows(values_only=True))
        self.assertEqual(rows[0][-1], "원고해시(sha256 앞 12자)")
        self.assertEqual(rows[1][-1], rk.reply_hash(LINK)[:12])


if __name__ == "__main__":
    unittest.main()
