#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
지식iN 답변 파이프라인 — 질문 수집 → 답변 초안(게이트) → PAMS 「지식인 심의」 접수 키트 → 승인 후 게시문.

    python scripts/kin_pipeline.py run [--n 3] [--dry-run]     # 수집 + 초안 + 키트(하루 상한 5건)
    python scripts/kin_pipeline.py approvals --xls <jidatexcel.xls>  # PAMS 「지식인&댓글 엑셀변환」으로 승인 수거
    python scripts/kin_pipeline.py approvals                    # (robert-os 감시기 승인 파일이 생기면) 그것으로
    python scripts/kin_pipeline.py sync-watcher [--dry-run]     # 감시기 승인 행 → ad_reviews(kin) 자동(review_no 중복 0)
    python scripts/kin_pipeline.py approve <키트id> --review-no 2026-10-1234 --from 2026-10-09 --to 2027-10-08 \\
                                           --notice-file <「심의내용」 돋보기 화면 복사본.txt>
    python scripts/kin_pipeline.py posted <키트id> --url <지식iN 답변 URL>
    python scripts/kin_pipeline.py weekly [--profile N] [--kakao N]   # 주간 성적표 한 줄

🔴 이 스크립트가 하지 않는 것
  · PAMS 「저장」「제출」 — 이 파일은 PAMS 에 접속하지 않는다. 「저장」 금지, 「제출」은 로버트만.
    폼 채우기는 robert-os pams_apply 몫이고, 그 전에 jumgumpyoji.jsp?nums=1 의 칸 이름 실측이 필요하다
    (docs/KIN-PIPELINE.md 「robert-os 할 일」). 그때까지 키트 txt 를 보고 사람이 옮겨 넣는다.
  · 네이버 로그인·게시 — 게시는 로버트가 최종본을 복사해 붙여 넣는다.
  · 링크 — 지식iN 답변에는 어떤 URL 도 넣지 않는다(kin-terms 게이트).

키트 위치: %USERPROFILE%\\Downloads\\PAMS접수\\_지식인\\  (🔴 하위 폴더다 — robert-os pams_apply 는 PAMS접수\\*.txt 만
    읽으므로 여기 두면 일반심의 폼으로 잘못 열리지 않는다.)
상태: %LOCALAPPDATA%\\SHIN\\kin_state.json — 키트별 상태(심의대기→승인→게시) · 하루 상한 · 소재 이력.
"""
import argparse
import hashlib
import json
import os
import re
import subprocess
import sys
import tempfile
from datetime import date, datetime, timedelta

import requests

HERE = os.path.dirname(os.path.abspath(__file__))
sys.path.insert(0, HERE)
try:
    sys.stdout.reconfigure(encoding="utf-8")
except Exception:
    pass

import kin_harvest  # noqa: E402
import pams_kit as kit  # noqa: E402

ROOT = kit.ROOT
STATE_DIR = kin_harvest.STATE_DIR
STATE_PATH = os.path.join(STATE_DIR, "kin_state.json")
KIT_DIR = os.path.join(kit.KIT_DIR, "_지식인")
WATCHER_JSON = os.environ.get("KIN_WATCHER", r"D:\robert-os\finance\state\pams_지식인행.json")
APPROVALS_JSON = os.environ.get("KIN_APPROVALS", r"D:\robert-os\finance\state\pams_kin_승인.json")
DAILY_MAX = 5
NOTE_MAX_BYTES = 100      # 특이사항(TB_GUMCHECK_LIST.CHSAYU) — 넘으면 ORA-12899 로 접수 실패(§6.4 스레드 주의 7)
REVIEW_NO_RX = re.compile(r"^\d{4}-(0[1-9]|1[0-2])-\d{4,5}$")   # src/lib/review-no.ts REVIEW_NO_RE 와 같다
DATE_RX = re.compile(r"^\d{4}-\d{2}-\d{2}$")
# PAMS 자동생성 필수안내사항의 자리표시자(매뉴얼 p5·§6.4 스레드 주의 8: 「프라임에셋 심의필 제0000호 (2023.00.00~2024.00.00)」)
PLACEHOLDER_RX = re.compile(r"제\s*0{3,5}\s*호\s*\(\s*\d{4}\.00\.00\s*~\s*\d{4}\.00\.00\s*\)")
# PAMS 지식인 폼 카테고리 라디오 6종(지식인_신청폼.png 실물 2026-10-09): 운전자보험 · 유병자보험 · 종신보험 ·
# 변액/연금보험 · 그외 기타(건강,화재,펫 보험 등) · 태아보험. 고르면 「답변내용」 아래 칸에 안내문구+필수안내사항 자동생성.
# 칸의 name/value 는 아직 재지 않았다 — 글자(라벨)만 적는다.
CATEGORY_ETC = "그외 기타(건강,화재,펫 보험 등)"
CATEGORY = {"간편심사": "유병자보험"}          # 나머지 소재는 전부 CATEGORY_ETC
POSTING_LOCATION = "네이버 지식인"             # 폼 기본값(실물). 고치지 않는다.
NOTICE_MARK = "※ 유의사항"                     # src/lib/compliance/kin-terms.ts KIN_NOTICE_MARK 와 같다


class KinError(Exception):
    pass


# ── 상태 ────────────────────────────────────────────────────
def load_state(path=STATE_PATH):
    try:
        return json.load(open(path, encoding="utf-8"))
    except (FileNotFoundError, json.JSONDecodeError):
        return {"kits": {}}


def save_state(state, path=STATE_PATH):
    os.makedirs(os.path.dirname(path), exist_ok=True)
    json.dump(state, open(path + ".tmp", "w", encoding="utf-8"), ensure_ascii=False, indent=1)
    os.replace(path + ".tmp", path)


# ── 순수 함수(테스트 대상) ─────────────────────────────────────
def pick(candidates, state, today, n):
    """오늘 더 만들 질문. 하루 상한 DAILY_MAX · 같은 소재 하루 1건 · 직전 날 마지막 소재와 연속 금지 · 이미 다룬 질문 제외."""
    kits = state.get("kits", {})
    today_kits = [k for k in kits.values() if k.get("day") == today]
    room = max(0, min(n, DAILY_MAX - len(today_kits)))
    used_topics = {k["topic"] for k in today_kits}
    prev = sorted((k for k in kits.values() if k.get("day", "") < today), key=lambda k: k.get("created", ""))
    if prev:
        used_topics.add(prev[-1]["topic"])
    done_urls = {k["question_url"] for k in kits.values()}
    out = []
    for c in candidates:
        if len(out) >= room:
            break
        if c["url"] in done_urls or c["topic"] in used_topics:
            continue
        out.append(c)
        used_topics.add(c["topic"])
    return out


def in_scope(candidates):
    """키트를 만들기 직전 한 번 더 — 보험 질문이 아니면 뺀다(KQ1: 「부동산 중개인 고지의무위반」이 심의까지 갔다).
    → (남길 후보, [(뺀 후보, 사유)]). 분류는 질문 URL 의 dirId 로 본다."""
    keep, off = [], []
    for c in candidates:
        m = kin_harvest.DIR_RX.search(c.get("url") or "")
        why = kin_harvest.scope_exclusion(c.get("title"), f"{c.get('title') or ''} {c.get('summary') or ''}",
                                          m.group(1) if m else "")
        (off.append((c, why)) if why else keep.append(c))
    return keep, off


def special_note(topic):
    """특이사항 칸 — 100바이트 이내. 넘으면 접수가 통째로 실패한다."""
    s = f"지식iN 답변 · {topic} · 링크 없음 · 증빙 해당 없음"
    if len(s.encode("utf-8")) > NOTE_MAX_BYTES:
        s = f"지식iN 답변 · {topic}"
    if len(s.encode("utf-8")) > NOTE_MAX_BYTES:
        raise KinError(f"특이사항이 {NOTE_MAX_BYTES}바이트를 넘는다 — {s}")
    return s


def kit_id(question_url, day):
    m = re.search(r"docId=(\d+)", question_url or "")
    return f"{day.replace('-', '')[4:]}_지식인_{m.group(1) if m else 'x'}"


def compose_answer(body, notice_block):
    return f"{body.strip()}\n\n{notice_block}" if notice_block else body.strip()


def body_part(answer):
    i = answer.find(NOTICE_MARK)
    return answer if i == -1 else answer[:i]


def normalize_review_no(raw):
    """src/lib/review-no.ts normalizeReviewNo 와 같은 규칙 — 번호만(2026-10-1234). 아니면 None."""
    if not isinstance(raw, str):
        return None
    s = re.sub(r"\s+", "", raw)
    s = re.sub(r"\(.*\)$", "", s)
    at = s.rfind("심의필")
    if at >= 0:
        s = s[at + 3:]
    s = re.sub(r"호$", "", re.sub(r"^제", "", s))
    return s if REVIEW_NO_RX.match(s) else None


def fill_notice(notice_text, review_no, f, t):
    """PAMS 자동생성 [안내문구]+[필수안내사항] → 실번호가 든 문구. 다른 글자는 한 글자도 바꾸지 않는다.

    · 자리표시자(「제0000호 (2023.00.00~2024.00.00)」)가 한 곳 있으면 거기에 실번호·실기간만 넣는다.
    · 승인 뒤 「심의내용」 화면·엑셀변환본은 이미 실번호가 들어 있다(매뉴얼 p15) → 그 번호가 맞으면 그대로 쓴다.
    · 둘 다 아니면 멈춘다 — 손으로 채우지 않는다."""
    if not REVIEW_NO_RX.match(review_no or ""):
        raise KinError(f"심의필 번호 형식이 다르다 — 「{review_no}」(번호만, 예: 2026-10-1234)")
    if not (DATE_RX.match(f or "") and DATE_RX.match(t or "")) or f > t:
        raise KinError(f"유효기간 형식·순서가 다르다 — {f}~{t}")
    hits = PLACEHOLDER_RX.findall(notice_text or "")
    if len(hits) == 1:
        return PLACEHOLDER_RX.sub(f"제{review_no}호 ({f.replace('-', '.')}~{t.replace('-', '.')})", notice_text, count=1)
    if not hits and re.search(rf"제\s*{re.escape(review_no)}\s*호", notice_text or ""):
        return notice_text
    raise KinError(f"필수안내사항에서 심의필 자리를 못 정했다(자리표시자 {len(hits)}곳, 실번호 없음) — 손으로 채우지 않고 멈춘다")


def _flat(s):
    return re.sub(r"\s+", "", s or "")


def compose_final(kit_rec, notice_text, review_no, f, t, today):
    """승인 후 게시문.

    · notice_text 가 PAMS 「심의내용」 화면(또는 엑셀변환 「답변내용」)의 **전문**이면 — 그 안에 우리 답변이 이미 있다 —
      그것이 게시문 원본이다(매뉴얼 p15 「해당 내용 복사하여 사용」). 답변이 글자 그대로(공백 무시) 들어 있는지만 보고 그대로 쓴다.
    · 자동생성 [안내문구]+[필수안내사항]만이면 = 심의받은 답변 그대로 + 빈 줄 + 그 문구(실번호).
    """
    answer = kit_rec["answer"]
    if hashlib.sha256(answer.encode("utf-8")).hexdigest() != kit_rec["answer_sha256"]:
        raise KinError("답변이 키트(심의본)와 다르다 — 한 글자라도 바뀌면 신규 심의")
    if not (f <= today <= t):
        raise KinError(f"유효기간 밖({f}~{t})")
    filled = fill_notice(notice_text, review_no, f, t).strip()
    body = _flat(body_part(answer))
    if body and body in _flat(filled):
        return filled
    if _flat(filled)[:40] == body[:40]:
        raise KinError("PAMS 문구가 답변으로 시작하는데 우리 답변과 글자가 다르다 — 심의본이 바뀌었는지 사람이 본다")
    return f"{answer}\n\n{filled}"


KIT_FIELDS = ("심의유형", "손보/생보", "광고형태", "게시위치", "카테고리", "특이사항", "파일첨부", "지식iN 질문", "원고해시")


def kit_text(rec):
    """PAMS 지식인 폼에 옮길 값. 형식 고정(SH4 2026-10-09, test_kin_pipeline 이 지킨다) — robert-os pams_apply 가 읽는다.

    · 1줄: `[PAMS 접수 문자열] 지식인 · <소재> · <키트id>`
    · 그다음 빈 줄까지 = 필드 블록. 줄마다 `<KIT_FIELDS 이름>: <값>` 하나, 순서 고정. **값 뒤에 설명을 붙이지 않는다**
      (본진 키트 「게시위치:」「규격:」 줄과 같은 규칙 — 설명은 아래 「안내」 블록).
    · `── PAMS 에 붙여 넣을 원고 ──` 다음 줄부터 `── 끝 ──` 앞 줄까지 = 「답변내용」 에디터에 넣을 원고(바이트 그대로).
    · 나머지(안내·질문 참고)는 사람용 — 기계는 읽지 않는다.
    폼 칸 순서: 카테고리를 **먼저** 고른다(자동생성 칸이 채워짐) → 원고. 「저장」·「제출」·자가점검은 기계 금지."""
    fields = {
        "심의유형": "지식인",
        "손보/생보": "손보",
        "광고형태": "바이럴(지식in)",
        "게시위치": POSTING_LOCATION,
        "카테고리": rec["category"],
        "특이사항": rec["note"],
        "파일첨부": "없음",
        "지식iN 질문": rec["question_url"],
        "원고해시": f"sha256 {rec['answer_sha256']}",
    }
    return "\n".join([
        f"[PAMS 접수 문자열] 지식인 · {rec['topic']} · {rec['id']}",
        *[f"{k}: {fields[k]}" for k in KIT_FIELDS],
        "",
        "── 안내(사람용) ──",
        "· 폼: jumgumpyoji.jsp?nums=1 — 광고형태 「바이럴(지식in)」·게시위치 「네이버 지식인」은 폼 기본값 그대로.",
        "· 카테고리를 먼저 고른다 → 「답변내용」 아래 칸에 안내문구+필수안내사항이 자동생성된다(지우지 않는다).",
        "· 원고는 텍스트만(이미지 없음). 자가점검은 사람이 확인 후 체크.",
        "· 파일첨부 없음 = 증빙 인용 없음(통계·의학정보·기관 귀속 문장 미사용).",
        "· 「지식iN 질문」은 답변을 달 곳이다(PAMS 칸 아님).",
        "",
        "── 질문 제목(참고) ──",
        rec["question_title"],
        "── 질문 본문(참고 — 검색 API 의 질문 조각만, 답변 조각 제외. PAMS 에 넣지 않는다: 질문자 정보 기재 금지 ⑤) ──",
        rec["question_summary"],
        "",
        "── PAMS 에 붙여 넣을 원고 ──",
        rec["answer"],
        "── 끝 ──",
        "",
        "🔴 PAMS 「저장」 금지 · 「제출」은 로버트만. 폼에 직접 고쳐 쓰지 않는다(사용 불가 표현이면 거절 + 작성 내용 삭제).",
        "승인 뒤(둘 중 하나):",
        "  python scripts/kin_pipeline.py approvals --xls <「지식인&댓글 엑셀변환」 jidatexcel.xls>",
        f"  python scripts/kin_pipeline.py approve {rec['id']} --review-no <번호> --from <YYYY-MM-DD> --to <YYYY-MM-DD> "
        "--notice-file <심의내용 돋보기 화면 복사본.txt>",
    ]) + "\n"


def weekly_line(state, today, profile=None, kakao=None):
    """주간 성적표 한 줄 — 지난 7일. 프로필 조회·카톡은 API 가 없어 사람이 넣는다(없으면 「미입력」).
    사이트 방문은 지식iN 답변에 링크가 없어 답변 단위로 귀속할 수 없다 → 「측정 불가」로 적는다(지어내지 않는다)."""
    start = (date.fromisoformat(today) - timedelta(days=7)).isoformat()
    recent = [k for k in state.get("kits", {}).values() if k.get("day", "") > start]
    approved = [k for k in recent if k.get("status") in ("승인", "게시")]
    posted = [k for k in recent if k.get("status") == "게시"]

    def fmt(v):
        return "미입력" if v is None else str(v)
    return (f"지식iN 주간({start[5:]}~{today[5:]}): 질문 {len(recent)} → 심의승인 {len(approved)} → 답변 게시 {len(posted)}"
            f" → 프로필 조회 {fmt(profile)} / 사이트 방문 측정 불가(링크 없음) / 카톡 {fmt(kakao)}")


# ── 게이트·초안(외부 호출) ────────────────────────────────────
def _tsx(script, payload):
    with tempfile.NamedTemporaryFile("w", suffix=".json", delete=False, encoding="utf-8") as fp:
        json.dump(payload, fp, ensure_ascii=False)
        tmp = fp.name
    try:
        r = subprocess.run(f'npx tsx {script} "{tmp}"', cwd=ROOT, shell=True, capture_output=True,
                           text=True, encoding="utf-8")
        if not r.stdout.strip():
            raise KinError(f"{script} 실행 실패 — {r.stderr.strip()[:300]}")
        return json.loads(r.stdout.strip().splitlines()[-1])
    finally:
        os.unlink(tmp)


def gate(items):
    """[{answer, question}] → [{pass, bodyLength, findings}]"""
    return _tsx("scripts/check_kin.mts --batch", items)


def notice_blocks(items):
    """[{body, question}] → [유의문구 블록] — 정본 자구는 brand.ts(kin-terms.kinNoticeBlock)가 만든다."""
    return _tsx("scripts/check_kin.mts --notice", items)


def draft(items):
    """[{id, title, body, feedback?}] → [{id, text}|{id, error}]"""
    return _tsx("scripts/kin_draft.mts", items)


# ── 틀 반복 금지(SH6 2026-10-09 관제탑) ─────────────────────────
# 견본 3건이 모두 「제 의견으로는」으로 열고 「담보 단위로 펼쳐 … 대조」로 닫았다. 하루 3~5건이 같은 틀이면
# 네이버 어뷰징·저품질 판정 위험 → ① 질문마다 서로 다른 여닫는 방식·의견 어구를 지정하고(factory-prompt KIN_*)
# ② 첫 문장·끝 문장을 최근 10건과 비교해 70% 이상 비슷하면 다시 쓴다 ③ 고정 어구는 하루 1회까지.
SIMILAR_MAX = 0.70
HISTORY_N = 10
_PROMPT_TS = os.path.join(ROOT, "src", "lib", "factory-prompt.ts")


def _ts_array(name, path=_PROMPT_TS):
    """factory-prompt.ts 의 문자열 배열 — 단일 출처(손으로 옮기지 않는다)."""
    ts = open(path, encoding="utf-8").read()
    m = re.search(rf"export const {name} = \[(.*?)\];", ts, re.S)
    if not m:
        raise KinError(f"{name} 를 factory-prompt.ts 에서 못 찾았다")
    return re.findall(r'"([^"]+)"', m.group(1))


def fixed_phrases():
    """하루 1회까지인 고정 어구 = 의견 귀속 어구 전부 + 옛 고정 맺음말."""
    return _ts_array("KIN_OPINION_PHRASES") + ["담보 단위로 펼쳐"]


def sentences(body):
    return [s for s in (x.strip() for x in re.split(r"(?<=[.?!])\s+|\n+", body or "")) if s]


def _norm(s):
    return re.sub(r"[^0-9A-Za-z가-힣]", "", s or "")


def similarity(a, b):
    from difflib import SequenceMatcher
    a, b = _norm(a), _norm(b)
    return SequenceMatcher(None, a, b).ratio() if a and b else 0.0


def variety_findings(answer, history, today, phrases):
    """history = [{text, day}] (최근 kin_answers + 같은 실행에서 통과한 것). 걸리면 [{rule, term, reason}]."""
    body = body_part(answer).strip()
    ss = sentences(body)
    if not ss:
        return []
    out = []
    recent = history[:HISTORY_N]
    for label, mine, idx in (("첫 문장", ss[0], 0), ("마지막 문장", ss[-1], -1)):
        for h in recent:
            hs = sentences(body_part(h["text"]).strip())
            if not hs:
                continue
            r = similarity(mine, hs[idx])
            if r >= SIMILAR_MAX:
                out.append({"rule": "repeat", "term": f"{label} {int(r * 100)}% 유사: {hs[idx][:30]}",
                            "reason": f"최근 {HISTORY_N}건과 {label}이 {int(SIMILAR_MAX * 100)}% 이상 비슷하다 — 다른 방식으로 열고 닫는다"})
                break
    used_today = " ".join(h["text"] for h in history if h.get("day") == today)
    for p in phrases:
        if p in body and p in used_today:
            out.append({"rule": "repeat", "term": f"고정 어구 하루 2회: {p}",
                        "reason": "같은 고정 어구는 하루 1회까지 — 지정된 다른 어구를 쓴다"})
    return out


# ── 질의 적합 게이트(KQ2 2026-10-19 관제탑) ─────────────────────
# 10/19 승인 8건 중 3건(4518 치아·4520 양성종양·4532 전동킥보드)이 첫 두 문장을 「판단을 가르는 기준은」
# 「먼저 보실 것은」「시점으로 나눠 보면」으로만 채워 질문자가 물은 것에 답하지 않았다(로버트: 「질문자 의도에 제대로 대답하냐」).
# 세 겹으로 본다 — ① 모델이 뽑은 「질문 핵심」이 있다 ② 모델 자체 채점이 「예」 ③ 첫 두 문장이 질문 핵심 낱말을 하나 이상
# 받고, 답을 미루는 말(INTENT_DEFER_RE)만 있고 직접 답 표지(INTENT_DIRECT_RE)가 없지 않다.
# 불합격 → 사유를 붙여 1회 다시 쓴다 → 또 불합격이면 후보에서 뺀다(INTENT_TRIES).
INTENT_TRIES = 2
INTENT_DEFER_RE = re.compile(r"판단(?:을|이)?\s*가르|걸려\s*있|여러\s*가지|몇\s*가지|기준은|기준이\s*됩니다|먼저\s*보실\s*것|"
                             r"짚(?:을|어\s*볼)\s*부분|시점으로\s*나눠|헷갈리시는\s*지점|순서(?:로|가)\s*(?:보|중요)")
INTENT_DIRECT_RE = re.compile(r"통상|일반적으로|보통|대개|대체로|대부분|경우가\s*많습니다|흔합니다|"
                              r"해당(?:합니다|하지\s*않|될\s*수|되는\s*경우|되지\s*않)|아닙니다|어렵습니다|"
                              r"볼\s*수\s*있습니다|보기\s*어렵|가능성이\s*(?:높|낮|있)")
_INTENT_STOP = {"질문", "질문자", "궁금", "문의", "경우", "여부", "어떻게", "무엇", "어떤", "되나요", "인가", "인지", "관련",
                "대한", "있는지", "없는지", "하는지", "되는지", "보험", "해당", "가능", "한가지"}


def _intent_terms(core):
    """질문 핵심의 내용어 — 2자 이상 한글 어절에서 조사·어미를 한 번 뗀다(kin-terms personalTerms 와 같은 방식)."""
    out = []
    for m in re.finditer(r"[가-힣A-Za-z]{2,}", core or ""):
        w = re.sub(r"(?:으로|에서|에게|까지|부터|이나|이면|인지|인가|이냐|냐|은|는|이|가|을|를|에|의|와|과|도|만|로)$", "", m.group(0))
        if len(w) >= 2 and w not in _INTENT_STOP and w not in out:
            out.append(w)
    return out


def first_two(body):
    first_para = next((p for p in re.split(r"\n\s*\n", (body or "").strip()) if p.strip()), "")
    return " ".join(sentences(first_para)[:2])


def intent_findings(core, direct, body):
    """[{rule:'intent', term, reason}] — 비면 합격. body = 유의문구 블록 앞 본문."""
    if not (core or "").strip():
        return [{"rule": "intent", "term": "질문 핵심 없음",
                 "reason": "머리 첫 줄 「질문 핵심: …」에 질문자가 실제로 묻는 한 가지를 쓴다"}]
    out, lead = [], first_two(body)
    if direct is not True:
        out.append({"rule": "intent", "term": "자체 채점 「아니오」",
                    "reason": f"첫 두 문장이 「{core}」에 직접 답하지 않는다 — 통상적인 기준을 구체적으로 먼저 말한다"})
    terms = _intent_terms(core)
    # 낱말 앞 2자만 맞아도 받은 것으로 친다(「고지」↔「고지의무」, 「치아보험」↔「치아」)
    if terms and not any(t in lead or t[:2] in lead for t in terms):
        out.append({"rule": "intent", "term": f"첫 두 문장에 질문 핵심 낱말 없음({'·'.join(terms[:4])})",
                    "reason": "첫 두 문장이 질문자가 물은 것을 다루지 않는다"})
    if INTENT_DEFER_RE.search(lead) and not INTENT_DIRECT_RE.search(lead):
        out.append({"rule": "intent", "term": f"답을 미루는 서두: {INTENT_DEFER_RE.search(lead).group(0)}",
                    "reason": "「판단 요소」 나열은 직접 답 뒤 보조로만 — 첫 두 문장에 「통상·일반적으로 ~」 식의 답을 둔다"})
    return out


# ── 질문 의도 — 검색 API 의 제목 + 질문 조각만(KQ4 2026-10-19 관제탑) ───────────
# 로버트(15:15): 「지식인은 니가 질문의 의도를 파악해서 답변을 맞게 제대로 작성해야지, 내가 그걸 왜 붙여넣어」.
# → 사람이 원문을 옮기는 경로(KQ3 kin_questions\<docId>.txt)는 쓰지 않는다.
# 🔴 질문 상세 페이지는 여전히 읽지 않는다 — kin.naver.com/robots.txt(KQ3 실측)가 ClaudeBot·anthropic-ai 등을 Disallow: / 로
#    막고 AI 학습·RAG 목적 봇 접근을 금지한다. 정식 경로 = 네이버 검색 오픈API(지식iN)의 title·description 뿐이고,
#    description 은 KQ2 필터(answerNo 없음 · 답변 표지 없음)를 통과한 질문 조각만 쓴다(kin_harvest.summarize).
#
# 의도 명확도 게이트 — 제목+조각에서 아래 셋이 **모두** 잡혀야 후보다. 못 잡으면 사유를 남기고 뺀다(질보다 양이 아니다).
#   ① 묻는 것: 물음 표지가 있는 한 문장(제목 우선)
#   ② 상품·담보 종류: 「치아보험」「실비」「배상책임」「부담보」처럼 무엇에 관한 질문인지(「보험」 단독은 안 친다)
#   ③ 질문자 상황 사실 1개 이상: 시점·단계·거절·청구 같은 사실(SITUATION_KINDS)
# 뽑은 것은 초안기에 「질문 의도」로 넘기고(kin_draft.mts), 채택 기준 게이트가 본문이 이것을 다뤘는지 잰다.
# 물음표 · 물음 어미 · 의문사, 또는 끝이 「~랍니다(들은 말)」「~되면/하면(조건만 던짐)」인 제목
ASK_RE = re.compile(r"[?？]"
                    r"|나요|까요|가요|인가|인지|는지|은지|을지|할지|될지|되나|될까|맞나|없나|있나|뭔가요|뭐죠|궁금|알려\s*주세요"
                    r"|어떻게|어떡|어쩌|왜|언제|얼마나|무엇|뭘|뭐가|가능한지|가능할까|받을\s*수\s*있|(?:게|건|것이)\s*맞"
                    r"|(?:랍니다|라는데요?|라고\s*하네요|되면|하면)\s*\.{0,3}\s*$")
PRODUCT_RE = re.compile(
    r"(?:실손|실비|치아|암|간병인?|치매|운전자|종신|정기|연금|저축|상해|질병|건강|화재|배상\s*책임|일상\s*생활\s*배상|어린이|태아|"
    r"유병자|간편\s*심사|간편|종합|생명|정액|수술비|입원비?|자동차|여행자|펫|변액|CI|GI)\s*보험|"
    r"실손|실비|진단비|수술비|입원\s*(?:비|일당)|일당|후유\s*장해|사망\s*보험금|배상\s*책임|휴업\s*손해|부담보|특약|보장성")
# 질문자 상황 종류 — (이름, 질문에서 잡는 식, 답변 본문이 그 상황을 다뤘다고 보는 식)
# 본문 식은 일반 낱말이다 — 질문의 구체 사정(검사 이름·다친 경위)을 되받으면 personal 게이트(kin-terms)가 막는다.
SITUATION_KINDS = [
    ("가입 시점", re.compile(r"(?:19|20)?\d{2}\s*년\s*(?:\d+\s*월)?|\d+\s*(?:개월|달|년)(?:\s*\d+\s*개월)?\s*(?:이\s*)?(?:전|후|째|만|지나|지났|넘)|"
                          r"가입\s*(?:전|후|직후|당시|하고|했는데|한\s*지)|보험\s*가입\s*후|작년|재작년|몇\s*년\s*전|오래\s*전"),
     re.compile(r"가입\s*(?:전|후|직후|당시|시점|한\s*뒤|하신\s*뒤|하기\s*전|하실\s*때|하셨)|시점|\d+\s*(?:년|개월)|기간")),
    ("가입 예정", re.compile(r"가입\s*(?:하려|하려고|할까|해야|예정|하는\s*게|하는게|할\s*수)|들려고|들어야|들까"),
     re.compile(r"가입\s*(?:전|하시기\s*전|하실\s*때|하려|을\s*앞두)|청약")),
    ("청구", re.compile(r"청구|보험금\s*(?:신청|받|지급)|지급\s*(?:요청|신청)"), re.compile(r"청구")),
    ("거절·부지급", re.compile(r"거절|부지급|지급\s*(?:불가|거부|안\s*(?:된|해|됨))|불가\s*(?:랍니다|하다|통보)|못\s*받"),
     re.compile(r"거절|부지급|지급(?:하지|되지)\s*않|지급\s*불가")),
    ("고지", re.compile(r"(?<![사최])고지|알릴\s*의무|알리지"), re.compile(r"(?<![사최])고지|알릴\s*의무|알리")),   # 「사고지만」 제외
    ("부담보", re.compile(r"부담보"), re.compile(r"부담보")),
    ("진단·발견", re.compile(r"(?<!무)진단(?!형)|발견|확인\s*되|판정|발병|생겼|생기|양성|종양"), re.compile(r"진단|발견|확인된|확인되")),
    ("사고", re.compile(r"사고|다쳤|다쳐|부상|넘어|낙상|부딪"), re.compile(r"사고|다치|부상")),
    ("검사·검진", re.compile(r"검사|검진|재검|엑스레이|촬영"), re.compile(r"검사|검진")),
    ("치료·수술·입원", re.compile(r"치료|수술|입원|통원|시술"), re.compile(r"치료|수술|입원|통원")),
    ("해지·갱신", re.compile(r"해지|해약|갱신|만기|실효"), re.compile(r"해지|해약|갱신|만기|실효")),
    ("여러 계약", re.compile(r"중복|여러\s*(?:개|건|군데)|두\s*(?:개|군데|곳)|겹|각각"), re.compile(r"중복|겹|여러\s*계약|계약마다|각각")),
    ("사망·수익자", re.compile(r"사망|수익자|상속"), re.compile(r"사망|수익자")),
    ("서류 발급", re.compile(r"발급|진단서|서류"), re.compile(r"발급|서류|진단서")),
    ("소득·생계", re.compile(r"생계|소득|월급|일을\s*못|일\s*못|휴업"), re.compile(r"휴업|소득|생계")),
]
INTENT_ASK_MAX = 80


def _sentences_of(text):
    return [s.strip() for s in re.split(r"(?<=[.?!？])\s+|\s*/\s*|\n+", text or "") if s.strip()]


def question_intent(title, summary):
    """검색 API 제목 + 질문 조각 → (의도 | None, 제외 사유). 의도 = {ask, products, situations:[이름]}.

    summary 는 kin_harvest.summarize 결과(KQ2 필터 통과 질문 조각만, 「 / 」로 이음)다."""
    title, summary = (title or "").strip(), (summary or "").strip()
    text = f"{title} {summary}"
    # 묻는 것 — 물음 문장 최대 2개(제목 · 조각 속 질문자의 물음). 긴 문장은 물음이 있는 뒷부분을 남긴다
    asks = [s if len(s) <= INTENT_ASK_MAX else "…" + s[-INTENT_ASK_MAX:]
            for s in [title] + _sentences_of(summary) if s and ASK_RE.search(s)]
    ask = " / ".join(asks[:2])
    products = []
    for m in PRODUCT_RE.finditer(text):
        p = re.sub(r"\s+", "", m.group(0))
        if p not in products:
            products.append(p)
    situations = [name for name, q_re, _ in SITUATION_KINDS if q_re.search(text)]
    missing = [label for label, ok in (("묻는 것", ask), ("상품·담보 종류", products), ("질문자 상황", situations)) if not ok]
    if missing:
        return None, "의도 불명확(" + "·".join(missing) + " 못 잡음)"
    return {"ask": ask, "products": products, "situations": situations}, ""


def split_by_intent(cands):
    """후보 → (의도가 잡힌 후보[intent 붙음], [(뺀 후보, 사유)])."""
    keep, drop = [], []
    for c in cands:
        intent, why = question_intent(c.get("title"), c.get("summary"))
        (keep.append({**c, "intent": intent}) if intent else drop.append((c, why)))
    return keep, drop


# ── 채택 기준 게이트(KQ3 → KQ4: 근거는 API 텍스트에서 뽑은 의도) — 질의 적합(KQ2) 확장 ──────────
# ① 상품·담보 + 질문자 상황 사실을 본문이 직접 다룬다 — 상품 1 + 상황 1 이상 = ADOPT_FACTS_MIN 개
# ② 첫 두 문장 직접 답 — intent_findings(KQ2)
# ③ 마지막 문단 = 다음 할 일(서류 + 요청·발급 같은 행동)
ADOPT_FACTS_MIN = 2
ACTION_DOC_RE = re.compile(r"청약서|약관|증권|진료\s*기록|진단서|소견서|영수증|진료비\s*(?:세부\s*)?(?:내역|명세)|서류|사본|"
                           r"확인서|안내문|통지서|거절\s*사유|심사\s*결과|고지\s*문항")
ACTION_VERB_RE = re.compile(r"요청|발급|받아\s*(?:두|보|놓)|챙겨|준비|신청|확인해\s*(?:보|두)|대조해|물어\s*보|받으시")
# 같은 담보를 가리키는 다른 이름 — 질문이 「실비」면 답의 「실손」도 받은 것으로 친다
PRODUCT_ALIASES = {"실비": ("실손", "실비"), "실비보험": ("실손", "실비"), "실손보험": ("실손", "실비")}


def product_covered(products, body):
    b = re.sub(r"\s+", "", body or "")
    for p in products or []:
        stem = re.sub(r"보험$", "", p)
        keys = PRODUCT_ALIASES.get(p, ()) + (p, stem if len(stem) >= 2 else p)
        if any(k in b for k in keys):
            return p
    return None


def covered_situations(situations, body):
    kinds = {name: b_re for name, _, b_re in SITUATION_KINDS}
    return [s for s in situations or [] if s in kinds and kinds[s].search(body or "")]


def last_paragraph(body):
    paras = [p.strip() for p in re.split(r"\n\s*\n", (body or "").strip()) if p.strip()]
    return paras[-1] if paras else ""


def adoption_findings(intent, body):
    """[{rule:'adopt', term, reason}] — 비면 합격. intent = question_intent 결과, body = 유의문구 블록 앞 본문."""
    out = []
    intent = intent or {}
    prod = product_covered(intent.get("products"), body)
    sits = covered_situations(intent.get("situations"), body)
    got = (1 if prod else 0) + len(sits)
    if not prod or not sits or got < ADOPT_FACTS_MIN:
        out.append({"rule": "adopt",
                    "term": f"질문자 상황 반영 {got}/{ADOPT_FACTS_MIN}(상품 {'·'.join(intent.get('products') or []) or '없음'}"
                            f"{' ✓' if prod else ' ✗'} · 상황 {'·'.join(intent.get('situations') or []) or '없음'}"
                            f" → 다룸 {'·'.join(sits) or '없음'})",
                    "reason": "질문 제목·조각에 적힌 상품·담보 종류와 질문자 상황(시점·단계·거절·청구 등)을 본문이 직접 다룬다"})
    lp = last_paragraph(body)
    if not (ACTION_DOC_RE.search(lp) and ACTION_VERB_RE.search(lp)):
        out.append({"rule": "adopt", "term": "마지막 문단에 다음 할 일 없음",
                    "reason": "마지막 문단은 질문자가 다음에 할 일 1~2가지(어떤 서류를 어디에 요청·발급하는지)로 끝낸다"})
    return out


def load_history(env, limit=HISTORY_N):
    """kin_answers 최근 limit 건(답변 초안·생성일 KST). 조회 실패는 빈 목록(같은 실행 안에서만 비교)."""
    try:
        url, h = _rest(env, "kin_answers")
        r = requests.get(url, params={"select": "answer_draft,created_at", "order": "created_at.desc", "limit": str(limit)},
                         headers=h, timeout=30)
        if r.status_code >= 300:
            return []
        out = []
        for row in r.json():
            day = ""
            try:
                day = datetime.fromisoformat(row["created_at"].replace("Z", "+00:00")).astimezone(kit.KST).date().isoformat()
            except Exception:  # noqa: BLE001
                pass
            if row.get("answer_draft"):
                out.append({"text": row["answer_draft"], "day": day})
        return out
    except Exception:  # noqa: BLE001
        return []


def style_plan(n, rng=None, n_open=None, n_close=None, n_opinion=None):
    """질문 n 건(+재작성분)에 서로 다른 여닫는 방식·의견 어구 번호 — 앞에서부터 꺼내 쓴다."""
    import random
    rng = rng or random.Random()
    n_open = n_open or len(_ts_array("KIN_OPENINGS"))
    n_close = n_close or len(_ts_array("KIN_CLOSINGS"))
    n_opinion = n_opinion or len(_ts_array("KIN_OPINION_PHRASES"))
    o, c, p = (rng.sample(range(k), k) for k in (n_open, n_close, n_opinion))
    return [{"open": o[i % n_open], "close": c[i % n_close], "opinion": p[i % n_opinion]} for i in range(n)]


def make_answers(cands, draft_fn=draft, gate_fn=gate, notice_fn=notice_blocks, tries=3, history=None, today=None,
                 phrases=None, styles=None, on_round=None):
    """후보 → [(후보, 답변 전문 | None, 막힌 사유)]. 게이트·반복 검사에 걸리면 사유를 붙여 처음부터 다시 쓴다(고쳐 쓰지 않는다).

    history = 최근 답변 [{text, day}](kin_answers) — 같은 실행에서 통과한 답변도 차례로 더해 비교한다.
    on_round(회차, 질문url, 답변, 틀, findings) — 회차별 판정 기록용(견본 보고)."""
    pending = {c["url"]: c for c in cands}
    feedback, result, intent_fail = {}, {}, {}
    history = list(history or [])
    today = today or datetime.now(kit.KST).date().isoformat()
    phrases = fixed_phrases() if phrases is None else phrases
    styles = styles if styles is not None else style_plan(len(cands) * tries)
    style_of, used = {}, 0
    for c in cands:   # 질문마다 서로 다른 틀, 재작성 때는 아직 안 쓴 틀
        style_of[c["url"]] = styles[used] if used < len(styles) else None
        used += 1
    for rnd in range(1, tries + 1):
        if not pending:
            break
        reqs = [{"id": u, "title": c["title"], "body": c["summary"], **({"feedback": feedback[u]} if u in feedback else {}),
                 **({"style": style_of[u]} if style_of.get(u) else {}),
                 **({"intent": c["intent"]} if c.get("intent") else {})}
                for u, c in pending.items()]
        drafts = {d["id"]: d for d in draft_fn(reqs)}
        for u in list(pending):
            if not drafts.get(u, {}).get("text"):
                result[u] = (None, f"초안 실패: {drafts.get(u, {}).get('error', '응답 없음')}")
                pending.pop(u)
        ok_items = [(u, drafts[u]["text"]) for u in pending]
        if not ok_items:
            break
        q = {u: pending[u]["title"] + "\n" + pending[u]["summary"] for u, _ in ok_items}
        blocks = notice_fn([{"body": t, "question": q[u]} for u, t in ok_items])
        answers = [(u, compose_answer(t, b)) for (u, t), b in zip(ok_items, blocks)]
        verdicts = gate_fn([{"answer": a, "question": q[u]} for u, a in answers])
        for (u, a), v in zip(answers, verdicts):
            findings = list(v["findings"]) if not v["pass"] else []
            if v["pass"]:
                findings = variety_findings(a, history, today, phrases)
            d = drafts[u]
            intent = intent_findings(d.get("core", ""), d.get("direct"), body_part(a))
            c = pending[u]
            if c.get("intent"):   # 채택 기준(KQ4) — cmd_run 은 의도가 잡힌 후보만 넘긴다
                intent += adoption_findings(c["intent"], body_part(a))
            findings += intent
            if on_round:
                on_round(rnd, u, a, style_of.get(u), findings)
            if intent:
                intent_fail[u] = intent_fail.get(u, 0) + 1
                if intent_fail[u] >= INTENT_TRIES:   # 재생성 1회 뒤에도 질문에 답하지 않으면 후보에서 뺀다
                    result[u] = (None, "질의 적합 불합격(채택 기준 포함, 재생성 1회 후): "
                                 + "; ".join(f["term"] for f in intent) + f" — 질문 핵심 「{d.get('core', '')}」")
                    pending.pop(u)
                    continue
            if not findings:
                result[u] = (a, "")
                history.insert(0, {"text": a, "day": today})
                pending.pop(u)
            else:
                if any(f["rule"] == "repeat" for f in findings) and used < len(styles):
                    style_of[u] = styles[used]   # 틀이 겹쳤으면 다음엔 다른 틀로
                    used += 1
                why = "; ".join(f"{f['rule']}: {f['term']}" for f in findings)
                v = {**v, "findings": findings}
                # 다시 쓰는 모델에는 규칙 사유까지 준다 — 「본문 946자」만으로는 줄일 목표(600~900자)를 모른다(SH5 견본 2)
                feedback[u] = "; ".join(f"{f['rule']}: {f['term']} ({f.get('reason', '')})" for f in v["findings"])
                result[u] = (None, f"게이트: {why}")
    return [(c, *result.get(c["url"], (None, "처리 안 됨"))) for c in cands]


# ── DB ──────────────────────────────────────────────────────
def _rest(env, table):
    url, h = kit._rest(env)
    return url.rsplit("/", 1)[0] + "/" + table, h


def save_draft_row(env, rec):
    """kin_answers 에 「초안」으로 — 어드민 지식iN 탭에 그대로 보인다. 실패해도 키트는 남긴다."""
    url, h = _rest(env, "kin_answers")
    r = requests.post(url, json={"question_url": rec["question_url"], "question_title": rec["question_title"],
                                 "question_body": rec["question_summary"] or rec["question_title"],
                                 "answer_draft": rec["answer"], "status": "초안"},
                      headers={**h, "Content-Type": "application/json", "Prefer": "return=representation"}, timeout=30)
    if r.status_code >= 300:
        return None, f"kin_answers 기록 실패 HTTP {r.status_code}"
    return r.json()[0]["id"], "kin_answers 초안"


def record_review(env, rec, review_no, f, t):
    """ad_reviews 에 channel='kin' 행. 🔴 article_id 가 NOT NULL·FK 라 지식인 답변(글이 아님)은 지금 넣을 수 없다
    (2026-10-09 OpenAPI 실측). sql/006_ad_reviews_kin.sql 적용 전에는 실패를 돌려주고 상태 파일에만 남긴다."""
    url, h = _rest(env, "ad_reviews")
    # channel='kin' = 지식iN 행의 유일한 식별 키(sql/006 체크·NOT_KIN 필터). review_type 은 PAMS 심의유형 — 기존 값 'jisikin'
    # (src/app/api/admin/ad-review/route.ts REVIEW_TYPES). 'kin' 은 review_type 체크에 없다.
    body = {"channel": "kin", "review_type": "jisikin", "ad_form": "바이럴(지식in)", "status": "approved",
            "posting_title": POSTING_LOCATION, "review_authority": "프라임에셋",
            "review_no": review_no, "review_from": f, "review_to": t,
            "reviewed_at": datetime.now(kit.KST).isoformat(),
            "notes": f"지식iN 답변 {rec['id']} · 질문 {rec['question_url']} · 원고 sha256 {rec['answer_sha256']}"}
    if not rec.get("kin_answer_id"):
        # sql/006 체크(ad_reviews_article_or_kin)가 kin 행에 kin_answer_id 를 요구한다 — 초안 행 기록이 실패한 키트
        return None, "ad_reviews 기록 못함 — kin_answers 초안 행이 없다(run 때 기록 실패). 상태 파일에만 남김"
    body["kin_answer_id"] = rec["kin_answer_id"]
    r = requests.post(url, json=body, headers={**h, "Content-Type": "application/json",
                                               "Prefer": "return=representation"}, timeout=30)
    if r.status_code >= 300:
        return None, f"ad_reviews 기록 못함(HTTP {r.status_code}) — sql/006 미적용이면 정상. 상태 파일에만 남김"
    return r.json()[0]["id"], "ad_reviews kin 행"


EXPIRY_WARN_DAYS = 30


def count_expiring(env, today, days=EXPIRY_WARN_DAYS):
    """ad_reviews channel=kin 승인 행 중 오늘~days일 안에 끝나는 것 수(§6.3 — 유효기간 만료는 미심의급 제재).
    지식iN 행은 article_id 가 NULL 이라 ad_reviews_expiring 뷰(premium_articles INNER JOIN)에 안 잡힌다 → 여기서 따로 센다."""
    url, h = _rest(env, "ad_reviews")
    end = (date.fromisoformat(today) + timedelta(days=days)).isoformat()
    r = requests.get(url, params=[("select", "id"), ("channel", "eq.kin"), ("status", "eq.approved"),
                                  ("review_to", f"gte.{today}"), ("review_to", f"lte.{end}")], headers=h, timeout=30)
    if r.status_code >= 300:
        raise KinError(f"HTTP {r.status_code}")
    return len(r.json())


def expiry_check(today, count_fn):
    """(건수 | None, 요약 줄). 조회 실패는 None — 수집·초안을 막지 않되 「실패」를 숨기지 않는다."""
    try:
        n = count_fn(today)
    except Exception as e:  # noqa: BLE001
        return None, f"지식iN 심의필 만료 {EXPIRY_WARN_DAYS}일 이내: 조회 실패({e})"
    return n, f"지식iN 심의필 만료 {EXPIRY_WARN_DAYS}일 이내 {n}건" + (" — PAMS 연장 신청 필요" if n else "")


def expiry_line(today, count_fn):
    return expiry_check(today, count_fn)[1]


def set_posted_url(env, rec):
    msgs = []
    if rec.get("ad_review_id"):
        url, h = _rest(env, "ad_reviews")
        r = requests.patch(url, params={"id": f"eq.{rec['ad_review_id']}"}, json={"posted_url": rec["posted_url"]},
                           headers={**h, "Content-Type": "application/json", "Prefer": "return=minimal"}, timeout=30)
        msgs.append(f"ad_reviews posted_url {'기록' if r.status_code < 300 else f'실패 {r.status_code}'}")
    if rec.get("kin_answer_id"):
        url, h = _rest(env, "kin_answers")
        r = requests.patch(url, params={"id": f"eq.{rec['kin_answer_id']}"},
                           json={"status": "게시완료", "posted_at": datetime.now(kit.KST).isoformat()},
                           headers={**h, "Content-Type": "application/json", "Prefer": "return=minimal"}, timeout=30)
        msgs.append(f"kin_answers 게시완료 {'기록' if r.status_code < 300 else f'실패 {r.status_code}'}")
    return msgs


# ── 엑셀변환(승인 수거) ──────────────────────────────────────
def _html_rows(text):
    """PAMS 엑셀변환(jidatexcel.xls) — JSP 가 내려주는 .xls 는 대개 HTML 표다. 셀 글자만 모은다."""
    from html.parser import HTMLParser
    rows, cur, cell, st = [], [], [], {"in": False}

    class P(HTMLParser):
        def handle_starttag(self, tag, attrs):
            if tag in ("td", "th"):
                st["in"] = True
                cell.clear()
            elif tag == "br" and st["in"]:
                cell.append("\n")
            elif tag == "tr":
                cur.clear()

        def handle_endtag(self, tag):
            if tag in ("td", "th"):
                st["in"] = False
                cur.append("".join(cell).strip())
            elif tag == "tr" and cur:
                rows.append(list(cur))

        def handle_data(self, data):
            if st["in"]:
                cell.append(data)

    P().feed(text)
    return rows


def read_export(path):
    """「지식인&댓글 엑셀변환」 → [{No, 신청일시, 변경일시, 심의필번호, 심의필일자, 광고유효기간, 답변내용}].
    양식(지식인댓글_엑셀변환_양식.png 실물 jidatexcel.xls): No · 신청일시 · 변경일시 · 심의필번호 · 심의필일자 · 광고유효기간 · 답변내용."""
    raw = open(path, "rb").read()
    if raw[:4] == b"PK\x03\x04":
        from openpyxl import load_workbook
        ws = load_workbook(path, read_only=True).active
        rows = [[("" if v is None else str(v)) for v in r] for r in ws.iter_rows(values_only=True)]
    elif b"<table" in raw[:8192].lower() or raw[:512].lstrip().startswith(b"<"):
        text = None
        for enc in ("utf-8-sig", "cp949"):
            try:
                text = raw.decode(enc)
                break
            except UnicodeDecodeError:
                continue
        if text is None:
            raise KinError("엑셀변환 파일 인코딩을 못 읽었다(utf-8·cp949 아님)")
        rows = _html_rows(text)
    else:
        try:
            import xlrd
        except ImportError:
            raise KinError("구형 .xls(BIFF) 다 — xlrd 가 없다. 엑셀에서 「다른 이름으로 저장 → .xlsx」 한 뒤 다시")
        sh = xlrd.open_workbook(path).sheet_by_index(0)
        rows = [[str(c.value) for c in sh.row(i)] for i in range(sh.nrows)]
    hdr_i = next((i for i, r in enumerate(rows) if "심의필번호" in [c.replace(" ", "") for c in r]), None)
    if hdr_i is None:
        raise KinError("엑셀변환 파일에서 「심의필번호」 머리줄을 못 찾았다")
    hdr = [c.replace(" ", "") for c in rows[hdr_i]]
    return [dict(zip(hdr, r)) for r in rows[hdr_i + 1:] if any(x.strip() for x in r)]


def _date(s):
    m = re.search(r"(\d{4})[-.](\d{1,2})[-.](\d{1,2})", s or "")
    return f"{m.group(1)}-{int(m.group(2)):02d}-{int(m.group(3)):02d}" if m else ""


def _period(row):
    """시작 = 심의필일자, 끝 = 광고유효기간(끝날짜만이거나 「시작 ~ 끝」)."""
    ds = re.findall(r"\d{4}[-.]\d{1,2}[-.]\d{1,2}", row.get("광고유효기간") or "")
    f = _date(row.get("심의필일자")) or (_date(ds[0]) if len(ds) > 1 else "")
    t = _date(ds[-1]) if ds else ""
    return f, t


def match_export(rows, kits):
    """엑셀 행 ↔ 우리 키트(심의대기) — 답변내용 안에 키트 답변 본문이 글자 그대로(공백 무시) 들어 있는 것만 짝."""
    out = []
    for row in rows:
        no = normalize_review_no(row.get("심의필번호") or "")
        if not no:
            continue   # 아직 심사중·반송
        flat = _flat(row.get("답변내용"))
        hits = [k for k in kits if k.get("status") == "심의대기" and _flat(body_part(k["answer"])) in flat]
        if len(hits) == 1:
            out.append((hits[0]["id"], no, *_period(row), row.get("답변내용") or ""))
    return out


# ── 명령 ────────────────────────────────────────────────────
def notify_fn(env, dry):
    if dry:
        return lambda t: print("[알림 생략 — dry-run]\n" + t)
    import pams_auto
    return lambda t: pams_auto.telegram(env, t)


def cmd_run(env, n, dry, harvest_fn=None, make_fn=None, notify=None, state_path=STATE_PATH,
            kit_dir=KIT_DIR, save_row=save_draft_row, now=None, expiring_fn=None, history_fn=None):
    now = now or datetime.now(kit.KST)
    today = now.date().isoformat()
    state = load_state(state_path)
    res = (harvest_fn or (lambda: kin_harvest.collect(env, now=now)))()
    scoped, off = in_scope(res["candidates"])
    # KQ4 — 제목+질문 조각으로 묻는 것·상품·상황을 못 잡은 질문은 초안을 만들지 않는다(일반론이 된다)
    clear, unclear = split_by_intent(scoped)
    chosen = pick(clear, state, today, n)
    expiring, exp_line = expiry_check(today, expiring_fn or (lambda d: count_expiring(env, d)))
    lines = [f"지식iN 수집: 질문 {res['seen_total']} · 후보 {len(res['candidates'])} · 제외 {len(res['excluded'])}"
             + (" · 첫 실행 기준선" if res.get("baseline") else "")
             + f" · 의도 명확 {len(clear)} · 의도 불명확 {len(unclear)} · 오늘 만들 것 {len(chosen)}", exp_line]
    lines += [f"⛔ 보험 무관 — {c['title'][:40]} ({why})" for c, why in off]
    lines += [f"🔍 [{c['topic']}] {c['title'][:40]} — {why}" for c, why in unclear]
    if make_fn is None:
        def make_fn(cs):
            return make_answers(cs, history=(history_fn or (lambda: load_history(env)))(), today=today)
    made = []
    for c, answer, why in (make_fn(chosen) if chosen else []):
        if not answer:
            lines.append(f"✖ [{c['topic']}] {c['title'][:40]} — {why[:300]}")
            continue
        rec = {"id": kit_id(c["url"], today), "day": today, "created": now.isoformat(), "topic": c["topic"],
               "question_url": c["url"], "question_title": c["title"], "question_summary": c["summary"],
               "answer": answer, "answer_sha256": hashlib.sha256(answer.encode("utf-8")).hexdigest(),
               "category": CATEGORY.get(c["topic"], CATEGORY_ETC),
               "note": special_note(c["topic"]), "status": "심의대기"}
        if not dry:
            rec["kin_answer_id"], _ = save_row(env, rec)
            os.makedirs(kit_dir, exist_ok=True)
            path = os.path.join(kit_dir, rec["id"] + ".txt")
            open(path, "w", encoding="utf-8", newline="\n").write(kit_text(rec))
            rec["kit"] = path
            state["kits"][rec["id"]] = rec
        made.append(rec)
        lines.append(f"✅ [{c['topic']}] {c['title'][:40]} → {rec['id']} (본문 {len(body_part(answer).strip())}자)")
    if not dry:
        save_state(state, state_path)
    msgs = []
    if made:
        msgs.append(f"🙋 지식인 심의 접수 대기 {len(made)}건 — PAMS 지식인 심의에 옮긴 뒤 「제출」은 로버트가(저장 금지)\n"
                    + "\n".join(f"· {r['id']} [{r['topic']}] {r['question_title'][:30]}" for r in made)
                    + f"\n키트: {kit_dir}")
    if expiring:   # 관제탑 결정(SH6) — N>0 이면 텔레그램에도 한 줄(§6.3 만료 = 미심의급 제재)
        msgs.append(f"⏰ {exp_line}")
    if msgs:
        (notify or notify_fn(env, dry))("\n\n".join(msgs))
    return lines, made


def cmd_approve(env, kid, review_no, f, t, notice_text, notify=None, state_path=STATE_PATH, record=record_review,
                today=None, dry=False, kit_dir=KIT_DIR):
    state = load_state(state_path)
    rec = state["kits"].get(kid)
    if not rec:
        raise KinError(f"키트 {kid} 가 상태 파일에 없다")
    review_no = normalize_review_no(review_no) or review_no   # 형식이 틀리면 fill_notice 가 멈춘다
    final = compose_final(rec, notice_text, review_no, f, t, today or date.today().isoformat())
    rec.update({"status": "승인", "review_no": review_no, "review_from": f, "review_to": t,
                "notice_text": notice_text, "final": final})
    if not dry:
        rec["ad_review_id"], rec["ad_review_msg"] = record(env, rec, review_no, f, t)
        out = os.path.join(kit_dir, kid + "_게시본.txt")
        os.makedirs(kit_dir, exist_ok=True)
        open(out, "w", encoding="utf-8", newline="\n").write(final + "\n")
        rec["final_path"] = out
        save_state(state, state_path)
    (notify or notify_fn(env, dry))(
        f"✅ 지식인 심의 승인 {kid} — 제{review_no}호. 아래를 그대로 복사해 지식iN 답변에 붙여 넣어 주세요"
        f"(게시 뒤 URL 을 알려 주시면 PAMS 게시위치·기록에 씁니다).\n질문: {rec['question_url']}\n\n{final}")
    return rec


def cmd_posted(env, kid, url, state_path=STATE_PATH, setter=set_posted_url):
    state = load_state(state_path)
    rec = state["kits"].get(kid)
    if not rec:
        raise KinError(f"키트 {kid} 가 상태 파일에 없다")
    if rec.get("status") != "승인":
        raise KinError(f"승인 전 키트다(상태 {rec.get('status')})")
    if not re.match(r"^https://(?:m\.)?kin\.naver\.com/", url or ""):
        raise KinError(f"지식iN 주소가 아니다 — {url}")
    rec.update({"status": "게시", "posted_url": url, "posted_at": datetime.now(kit.KST).isoformat()})
    msgs = setter(env, rec)
    save_state(state, state_path)
    return msgs


def cmd_approvals(env, xls=None, path=APPROVALS_JSON, state_path=STATE_PATH, notify=None, record=record_review,
                  today=None, kit_dir=KIT_DIR):
    """승인 수거 — ① PAMS 「지식인&댓글 엑셀변환」 파일(--xls) ② robert-os 감시기 승인 파일(계약은 docs/KIN-PIPELINE.md)."""
    state = load_state(state_path)
    if xls:
        items = match_export(read_export(xls), list(state["kits"].values()))
    elif os.path.exists(path):
        items = [(kid, a.get("심의필번호", ""), _date(a.get("시작")), _date(a.get("종료")), a.get("답변내용", ""))
                 for kid, a in (json.load(open(path, encoding="utf-8")) or {}).items()]
    else:
        return ["승인 파일 없음 — --xls 로 엑셀변환본을 주거나 approve 를 손으로"]
    out = []
    for kid, no, f, t, text in items:
        rec = state["kits"].get(kid)
        if not rec or rec.get("status") != "심의대기":
            continue
        try:
            cmd_approve(env, kid, no, f, t, text, notify=notify, state_path=state_path, record=record,
                        today=today, kit_dir=kit_dir)
            out.append(f"승인 처리 {kid} — 제{normalize_review_no(no) or no}호")
        except KinError as e:
            out.append(f"승인 처리 실패 {kid}: {e}")
    return out or ["새 승인 없음"]


def existing_review_nos(env):
    """ad_reviews channel=kin 에 이미 있는 review_no 집합 — 중복 삽입 방지의 기준."""
    url, h = _rest(env, "ad_reviews")
    r = requests.get(url, params=[("select", "review_no"), ("channel", "eq.kin")], headers=h, timeout=30)
    if r.status_code >= 300:
        raise KinError(f"ad_reviews 조회 실패 HTTP {r.status_code}")
    return {x["review_no"] for x in r.json() if x.get("review_no")}


def mark_answer_posted(env, kin_answer_id):
    """kin_answers.status 허용값은 초안·게시완료·채택됨·폐기 뿐(kin_answers_status_check) — 승인은 올릴 값이 없고 게시만 올린다."""
    url, h = _rest(env, "kin_answers")
    r = requests.patch(url, params={"id": f"eq.{kin_answer_id}"},
                       json={"status": "게시완료", "posted_at": datetime.now(kit.KST).isoformat()},
                       headers={**h, "Content-Type": "application/json", "Prefer": "return=minimal"}, timeout=30)
    return r.status_code < 300


def cmd_sync_watcher(env, path=WATCHER_JSON, state_path=STATE_PATH, record=record_review, existing_fn=None,
                     mark_posted=mark_answer_posted, dry=False):
    """감시기(robert-os pams_지식인행.json) 승인 행 → ad_reviews(kin) 자동 기록. review_no 가 이미 있으면 건너뛴다(중복 0).
    게시위치등록됨 이면 kin_answers 를 게시완료로. 사람이 approvals/approve 를 돌릴 필요가 없다."""
    if not os.path.exists(path):
        return ["감시기 파일 없음"]
    rows = (json.load(open(path, encoding="utf-8")) or {}).get("행", [])
    state = load_state(state_path)
    have = (existing_fn or (lambda: existing_review_nos(env)))()
    out, changed = [], False
    for row in rows:
        no = normalize_review_no(row.get("심의번호") or "")
        if not no:
            continue   # 심사중·반송
        kid = row.get("키트id") or row.get("키트")
        rec = state["kits"].get(kid)
        if no in have:
            continue
        if not rec:
            out.append(f"건너뜀 {no}: 키트 {kid} 가 상태 파일에 없다")
            continue
        f, t = _date(row.get("심의필일자")), _date(row.get("유효기간"))
        if dry:
            out.append(f"[dry-run] 기록 예정 {no} ← {kid}")
            continue
        ad_id, msg = record(env, rec, no, f, t)
        if not ad_id:
            out.append(f"기록 실패 {no} ← {kid}: {msg}")
            continue
        have.add(no)
        rec.update({"review_no": no, "review_from": f, "review_to": t, "ad_review_id": ad_id,
                    "status": "게시" if row.get("게시위치등록됨") else "승인"})
        if row.get("게시위치등록됨") and rec.get("kin_answer_id"):
            mark_posted(env, rec["kin_answer_id"])
        changed = True
        out.append(f"기록 {no} ← {kid} ({rec['status']})")
    if changed:
        save_state(state, state_path)
    return out or ["새 승인 없음"]


def main():
    ap =argparse.ArgumentParser(description="지식iN 답변 파이프라인")
    sub = ap.add_subparsers(dest="cmd", required=True)
    r = sub.add_parser("run")
    r.add_argument("--n", type=int, default=3)
    r.add_argument("--dry-run", action="store_true")
    a = sub.add_parser("approve")
    a.add_argument("kit")
    a.add_argument("--review-no", required=True)
    a.add_argument("--from", dest="f", required=True)
    a.add_argument("--to", dest="t", required=True)
    a.add_argument("--notice-file", required=True)
    a.add_argument("--dry-run", action="store_true")
    ap_ = sub.add_parser("approvals")
    ap_.add_argument("--xls", help="PAMS 「지식인&댓글 엑셀변환」 내려받은 파일")
    sw = sub.add_parser("sync-watcher")
    sw.add_argument("--dry-run", action="store_true")
    p = sub.add_parser("posted")
    p.add_argument("kit")
    p.add_argument("--url", required=True)
    w = sub.add_parser("weekly")
    w.add_argument("--profile", type=int)
    w.add_argument("--kakao", type=int)
    args = ap.parse_args()
    env = kit.load_env()
    try:
        if args.cmd == "run":
            lines, _ = cmd_run(env, args.n, args.dry_run)
            print("\n".join(lines))
        elif args.cmd == "approve":
            notice = open(args.notice_file, encoding="utf-8-sig").read()
            rec = cmd_approve(env, args.kit, args.review_no, args.f, args.t, notice, dry=args.dry_run)
            print(rec.get("final_path") or rec["final"])
        elif args.cmd == "approvals":
            print("\n".join(cmd_approvals(env, xls=args.xls)))
        elif args.cmd == "sync-watcher":
            print("\n".join(cmd_sync_watcher(env, dry=args.dry_run)))
        elif args.cmd == "posted":
            print("\n".join(cmd_posted(env, args.kit, args.url)))
        elif args.cmd == "weekly":
            print(weekly_line(load_state(), date.today().isoformat(), args.profile, args.kakao))
    except KinError as e:
        print(f"❌ {e}", file=sys.stderr)
        sys.exit(1)


if __name__ == "__main__":
    main()
